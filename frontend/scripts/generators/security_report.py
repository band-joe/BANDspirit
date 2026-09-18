from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from datetime import datetime

doc = Document()

# Page setup
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

# -- Custom styles --
def make_style(name, stype, font_name, size, bold=False, color=None, space_before=0, space_after=0, alignment=None):
    s = doc.styles.add_style(name, stype)
    s.font.name = font_name
    s.font.size = Pt(size)
    s.font.bold = bold
    if color:
        s.font.color.rgb = RGBColor(*color)
    s.paragraph_format.space_before = Pt(space_before)
    s.paragraph_format.space_after = Pt(space_after)
    if alignment is not None:
        s.paragraph_format.alignment = alignment
    return s

make_style('CoverTitle', WD_STYLE_TYPE.PARAGRAPH, 'Calibri', 28, True, (0x1a, 0x23, 0x7e), 120, 6, WD_ALIGN_PARAGRAPH.CENTER)
make_style('CoverSubtitle', WD_STYLE_TYPE.PARAGRAPH, 'Calibri', 14, False, (0x55, 0x55, 0x55), 6, 6, WD_ALIGN_PARAGRAPH.CENTER)
make_style('SectionHead', WD_STYLE_TYPE.PARAGRAPH, 'Calibri', 18, True, (0x1a, 0x23, 0x7e), 24, 12)
make_style('SubHead', WD_STYLE_TYPE.PARAGRAPH, 'Calibri', 13, True, (0x0d, 0x47, 0xa1), 16, 8)
make_style('BodyText', WD_STYLE_TYPE.PARAGRAPH, 'Calibri', 10.5, False, (0x33, 0x33, 0x33), 4, 4)
make_style('FindingTitle', WD_STYLE_TYPE.PARAGRAPH, 'Calibri', 11, True, (0x33, 0x33, 0x33), 10, 4)
make_style('SmallNote', WD_STYLE_TYPE.PARAGRAPH, 'Calibri', 9, False, (0x77, 0x77, 0x77), 2, 2, WD_ALIGN_PARAGRAPH.CENTER)

# ========== COVER PAGE ==========
for _ in range(3):
    doc.add_paragraph('', style='BodyText')

doc.add_paragraph('Sicherheitsaudit', style='CoverTitle')
doc.add_paragraph('BANDspirit – Klientenmanagement-System', style='CoverSubtitle')
doc.add_paragraph('', style='BodyText')
doc.add_paragraph(f'Erstellt: {datetime.now().strftime("%d.%m.%Y")}', style='CoverSubtitle')
doc.add_paragraph('Version 1.2 – Quellcode-Review', style='CoverSubtitle')
doc.add_paragraph('', style='BodyText')
doc.add_paragraph('Klassifizierung: VERTRAULICH', style='CoverSubtitle')

doc.add_page_break()

# ========== 1. ZUSAMMENFASSUNG ==========
doc.add_paragraph('1. Zusammenfassung', style='SectionHead')
p = doc.add_paragraph(style='BodyText')
p.add_run('Gegenstand: ').bold = True
p.add_run('Statische Quellcode-Analyse des BANDspirit-Klientenmanagement-Systems (Version 1.2) mit 11 Modulen, 28+ API-Routen und rollenbasierter Zugriffskontrolle (RBAC).')

p = doc.add_paragraph(style='BodyText')
p.add_run('Methodik: ').bold = True
p.add_run('Manuelle Code-Review aller sicherheitsrelevanten Bereiche: Authentifizierung, Autorisierung, Input-Validierung, Datenbankzugriffe, Client-Side-Sicherheit, HTTP-Header und Konfiguration.')

p = doc.add_paragraph(style='BodyText')
p.add_run('Gesamtbewertung: ').bold = True
p.add_run('Das System implementiert solide Sicherheitsgrundlagen (bcrypt-Hashing, JWT-Sessions, RBAC, Prisma ORM). Es wurden keine kritischen Schwachstellen (Severity: KRITISCH) identifiziert. Es bestehen jedoch mehrere Befunde mittlerer und niedriger Priorität, die zur Härtung des Systems adressiert werden sollten.')

# Summary table
doc.add_paragraph('', style='BodyText')
t = doc.add_table(rows=5, cols=2, style='Table Grid')
t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = t.rows[0].cells
hdr[0].text = 'Schweregrad'
hdr[1].text = 'Anzahl Befunde'
for c in hdr:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), '1a237e')
    shd.set(qn('w:val'), 'clear')
    c._tc.get_or_add_tcPr().append(shd)

data = [('KRITISCH', '0'), ('HOCH', '2'), ('MITTEL', '4'), ('NIEDRIG', '3')]
for i, (sev, count) in enumerate(data):
    t.rows[i+1].cells[0].text = sev
    t.rows[i+1].cells[1].text = count

# ========== 2. POSITIVE BEFUNDE ==========
doc.add_page_break()
doc.add_paragraph('2. Positive Sicherheitsbefunde', style='SectionHead')
doc.add_paragraph('Folgende Sicherheitsmaßnahmen sind korrekt implementiert:', style='BodyText')

positives = [
    ('Passwort-Hashing', 'bcrypt mit 10 Salt-Runden in allen relevanten Endpunkten (Signup, User-CRUD, Login). Passwörter werden nie im Klartext gespeichert oder an Clients zurückgegeben.'),
    ('JWT-basierte Sessions', 'NextAuth.js mit JWT-Strategie, NEXTAUTH_SECRET korrekt konfiguriert. Kein Session-Token-Leak möglich.'),
    ('RBAC durchgängig', 'Alle 28+ API-Routen prüfen Berechtigungen via requirePermission(). 5 Rollen mit granularen Berechtigungen (35+ Permissions) sind definiert.'),
    ('Prisma ORM (kein Raw SQL)', 'Ausschließlich Prisma-Client wird verwendet. Kein $queryRaw oder $executeRaw gefunden – SQL-Injection ist ausgeschlossen.'),
    ('Keine XSS-Vektoren', 'Kein dangerouslySetInnerHTML, eval() oder Function() im gesamten Frontend-Code. React\'s automatisches Escaping schützt die Ausgabe.'),
    ('Secrets-Management', 'Alle sensiblen Werte (DATABASE_URL, NEXTAUTH_SECRET, ABACUSAI_API_KEY) in .env. Keine NEXT_PUBLIC_-Variablen mit Secrets. Keine hartkodierten Schlüssel im Code.'),
    ('Passwort aus API-Responses ausgeschlossen', 'User-API-Routen verwenden explizite select-Klauseln, die das password-Feld ausschließen. Kein Passwort-Leak über die API.'),
    ('Soft-Delete für Benutzer', 'Benutzer werden deaktiviert statt gelöscht – bewahrt Referenzintegrität und Audit-Trail.'),
    ('Self-Protection', 'Admins können eigene Rolle/Status nicht ändern und sich nicht selbst deaktivieren.'),
    ('Middleware-geschützte Routen', 'NextAuth-Middleware schützt alle Routen außer explizit öffentlichen Pfaden (Login, Signup, Auth-API).'),
]

for title, desc in positives:
    doc.add_paragraph(f'✓  {title}', style='FindingTitle')
    doc.add_paragraph(desc, style='BodyText')

# ========== 3. BEFUNDE ==========
doc.add_page_break()
doc.add_paragraph('3. Sicherheitsbefunde & Maßnahmen', style='SectionHead')

findings = [
    {
        'id': 'SEC-001',
        'severity': 'HOCH',
        'title': 'Offener Signup-Endpunkt ohne Autorisierung',
        'desc': 'Der Endpunkt /api/signup ist öffentlich zugänglich und ermöglicht es jedem, ein neues Benutzerkonto mit der Rolle FALLMANAGER anzulegen. Es gibt keine Einschränkung (Einladungs-Token, Admin-Freigabe, CAPTCHA).',
        'risk': 'Unbefugte können sich selbst registrieren und erhalten sofort Zugriff auf Klientendaten (klient:read, klient:create, klient:update, intake:*, bericht:*, kontakt:*, meeting:*).',
        'action': 'Option A (empfohlen): Signup nur über Admin-Einladung ermöglichen – Endpunkt mit requirePermission(\'user:create\') schützen.\nOption B: Signup deaktivieren und Benutzer ausschließlich über die Benutzerverwaltung anlegen.\nOption C: Wenn Selbstregistrierung gewünscht, CAPTCHA + E-Mail-Verifikation + Admin-Freischaltung implementieren.',
    },
    {
        'id': 'SEC-002',
        'severity': 'HOCH',
        'title': 'Kein Rate-Limiting auf Authentifizierungs-Endpunkten',
        'desc': 'Die Endpunkte /api/auth/login, /api/signup und /api/auth/[...nextauth] haben kein Rate-Limiting. Brute-Force-Angriffe auf Passwörter sind unbegrenzt möglich.',
        'risk': 'Angreifer können automatisiert Passwörter durchprobieren. Bei schwachen Passwörtern (z.B. "Test123!") ist ein erfolgreicher Angriff wahrscheinlich.',
        'action': 'Rate-Limiting in der Middleware implementieren (z.B. max. 5 Login-Versuche pro Minute pro IP). Nach 10 Fehlversuchen: temporäre Kontosperre (15 Min.). Logging fehlgeschlagener Anmeldeversuche für Monitoring.',
    },
    {
        'id': 'SEC-003',
        'severity': 'MITTEL',
        'title': 'Keine Passwort-Komplexitätsanforderungen',
        'desc': 'Weder /api/signup noch /api/users (POST/PUT) validieren die Passwortstärke. Jedes beliebige Passwort (auch "a" oder "123") wird akzeptiert.',
        'risk': 'Benutzer können triviale Passwörter setzen, die leicht erraten werden können.',
        'action': 'Zod-Schema für Passwortvalidierung: Minimum 8 Zeichen, mindestens 1 Großbuchstabe, 1 Kleinbuchstabe, 1 Zahl, 1 Sonderzeichen. Validierung sowohl im Signup- als auch im User-CRUD-Endpunkt.',
    },
    {
        'id': 'SEC-004',
        'severity': 'MITTEL',
        'title': 'Fehlende Input-Validierung in mehreren API-Routen',
        'desc': 'Nur /api/klienten und /api/arbeitsplaetze verwenden Zod-Schema-Validierung. Alle anderen Endpunkte (Kontakte, Meetings, Meeting-Typen, Berichte, Intakes, Berufsbilder) validieren nur auf Pflichtfeld-Ebene ohne Typprüfung, Längenbeschränkung oder Format-Validierung.',
        'risk': 'Überlange Strings, unerwartete Datentypen oder ungültige Formate könnten gespeichert werden. Prisma schützt vor SQL-Injection, aber Datenintegrität und Speicher-Effizienz sind gefährdet.',
        'action': 'Zod-Schemas für alle Entitäten erstellen (Kontakt, Meeting, MeetingTyp, Bericht, Intake, Berufsbild). String-Felder mit max-Länge versehen (z.B. titel max. 200, notizen max. 10000). E-Mail-Format, Datums-Format und numerische Bereiche validieren.',
    },
    {
        'id': 'SEC-005',
        'severity': 'MITTEL',
        'title': 'Keine Security-HTTP-Header konfiguriert',
        'desc': 'Die Anwendung setzt keine zusätzlichen Sicherheits-Header: kein Content-Security-Policy, kein X-Content-Type-Options, kein Strict-Transport-Security, kein Referrer-Policy, kein Permissions-Policy.',
        'risk': 'Erhöhte Anfälligkeit für Clickjacking (teilweise), MIME-Sniffing, und fehlende HSTS-Protection. CSP würde als zusätzliche XSS-Schutzschicht dienen.',
        'action': 'Security-Header in next.config.js unter headers() hinzufügen:\n- X-Content-Type-Options: nosniff\n- Referrer-Policy: strict-origin-when-cross-origin\n- Permissions-Policy: camera=(), microphone=(self), geolocation=()\n- Strict-Transport-Security: max-age=31536000; includeSubDomains\nHinweis: X-Frame-Options darf NICHT gesetzt werden (App wird im iframe angezeigt).',
    },
    {
        'id': 'SEC-006',
        'severity': 'MITTEL',
        'title': 'JWT-Token enthält keine Ablaufzeit-Konfiguration',
        'desc': 'Die NextAuth-Konfiguration definiert keine explizite maxAge für JWT-Tokens. Es wird der Standard von 30 Tagen verwendet.',
        'risk': 'Kompromittierte Tokens sind 30 Tage gültig. Für eine Anwendung mit sensiblen Sozialdaten ist das ein unangemessen langer Zeitraum.',
        'action': 'JWT-maxAge auf 8 Stunden setzen (Schichtwechsel-Logik). Session-Refresh mit kürzerer Frist. Explizit in authOptions konfigurieren:\nsession: { strategy: "jwt", maxAge: 8 * 60 * 60 }',
    },
    {
        'id': 'SEC-007',
        'severity': 'NIEDRIG',
        'title': 'Fehlende Audit-Protokollierung für sensible Operationen',
        'desc': 'Audit-Logging existiert nur für den Abrechnungs-Workflow. Login-Versuche (erfolgreich/fehlgeschlagen), Benutzeränderungen, Klientendaten-Zugriffe und Meeting-Operationen werden nicht protokolliert.',
        'risk': 'Bei Sicherheitsvorfällen können Zugriffe und Änderungen nicht nachvollzogen werden. Für Einrichtungen im Sozialbereich bestehen möglicherweise gesetzliche Protokollierungspflichten.',
        'action': 'Audit-Logging erweitern:\n- Login-Versuche (Erfolg/Fehlschlag mit IP und Timestamp)\n- CRUD-Operationen auf Klientendaten\n- Benutzeränderungen (Rollenwechsel, Aktivierung/Deaktivierung)\n- Meeting-Aufnahmen und KI-Zusammenfassungen\n- Export-Operationen (DATEV-CSV)',
    },
    {
        'id': 'SEC-008',
        'severity': 'NIEDRIG',
        'title': 'Keine CORS-Einschränkung konfiguriert',
        'desc': 'Keine explizite CORS-Policy in next.config.js oder Middleware. Next.js verwendet standardmäßig Same-Origin, aber für API-Routen ist keine explizite Einschränkung definiert.',
        'risk': 'Potenziell können Anfragen von beliebigen Ursprüngen an die API gesendet werden, falls die Standard-Policy geändert wird.',
        'action': 'Explizite CORS-Policy in next.config.js definieren, die nur den eigenen Origin erlaubt. Für API-Routen: Access-Control-Allow-Origin auf die Deployment-Domain beschränken.',
    },
    {
        'id': 'SEC-009',
        'severity': 'NIEDRIG',
        'title': 'Detaillierte Fehlermeldungen in Konsolen-Logs',
        'desc': 'API-Routen loggen vollständige Fehlerobjekte via console.error(). In Produktionsumgebungen könnten diese Stack-Traces und interne Details enthalten.',
        'risk': 'Gering – serverseitige Logs sind nicht direkt für Endbenutzer einsichtbar. Jedoch könnte bei Fehlkonfigurationen die Detail-Ebene zu hoch sein.',
        'action': 'Strukturiertes Logging mit Log-Leveln einführen. In Produktion nur Error-Code und -Message loggen, keine vollständigen Stack-Traces. Erwägen eines Log-Aggregation-Dienstes.',
    },
]

for f in findings:
    sev_color = {'KRITISCH': 'FF0000', 'HOCH': 'E65100', 'MITTEL': 'F9A825', 'NIEDRIG': '2E7D32'}
    
    p = doc.add_paragraph(style='FindingTitle')
    run = p.add_run(f'[{f["severity"]}] ')
    run.font.color.rgb = RGBColor.from_string(sev_color.get(f['severity'], '333333'))
    run.bold = True
    p.add_run(f'{f["id"]}: {f["title"]}')
    
    p = doc.add_paragraph(style='BodyText')
    p.add_run('Beschreibung: ').bold = True
    p.add_run(f['desc'])
    
    p = doc.add_paragraph(style='BodyText')
    p.add_run('Risiko: ').bold = True
    p.add_run(f['risk'])
    
    p = doc.add_paragraph(style='BodyText')
    p.add_run('Maßnahme: ').bold = True
    p.add_run(f['action'])
    
    doc.add_paragraph('', style='BodyText')

# ========== 4. PRÜFMATRIX ==========
doc.add_page_break()
doc.add_paragraph('4. Prüfmatrix – Geprüfte Sicherheitsbereiche', style='SectionHead')

matrix_data = [
    ('Authentifizierung', 'bcrypt-Hashing, JWT-Sessions, aktiv-Flag-Prüfung', '✓ Bestanden'),
    ('Autorisierung (RBAC)', '28+ Routen geprüft, requirePermission() durchgängig', '✓ Bestanden'),
    ('SQL-Injection', 'Kein Raw SQL, ausschließlich Prisma ORM', '✓ Bestanden'),
    ('XSS (Cross-Site-Scripting)', 'Kein dangerouslySetInnerHTML, eval(), innerHTML', '✓ Bestanden'),
    ('CSRF-Schutz', 'JWT in httpOnly-Cookie via NextAuth', '✓ Bestanden'),
    ('Secrets-Management', 'Alle Keys in .env, keine Client-Exposition', '✓ Bestanden'),
    ('Passwort-Speicherung', 'bcrypt, nie im Klartext, nie in API-Responses', '✓ Bestanden'),
    ('Input-Validierung', 'Teilweise (nur Klienten/Arbeitsplätze mit Zod)', '⚠ Teilweise'),
    ('Rate-Limiting', 'Nicht implementiert', '✗ Fehlend'),
    ('Security-Header', 'Nicht konfiguriert', '✗ Fehlend'),
    ('Audit-Logging', 'Nur Abrechnungs-Workflow', '⚠ Teilweise'),
    ('Session-Timeout', 'Standard 30 Tage, nicht angepasst', '⚠ Anpassen'),
    ('Signup-Kontrolle', 'Öffentlich ohne Einschränkung', '✗ Fehlend'),
    ('CORS-Policy', 'Standard (Same-Origin), nicht explizit', '⚠ Empfohlen'),
    ('Error-Handling', 'console.error mit vollen Objekten', '⚠ Verbessern'),
]

t = doc.add_table(rows=len(matrix_data)+1, cols=3, style='Table Grid')
t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = t.rows[0].cells
for i, h in enumerate(['Prüfbereich', 'Ergebnis', 'Status']):
    hdr[i].text = h
    for p in hdr[i].paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), '1a237e')
    shd.set(qn('w:val'), 'clear')
    hdr[i]._tc.get_or_add_tcPr().append(shd)

for i, (area, result, status) in enumerate(matrix_data):
    row = t.rows[i+1]
    row.cells[0].text = area
    row.cells[1].text = result
    row.cells[2].text = status

# ========== 5. PRIORISIERTE MAßNAHMEN ==========
doc.add_page_break()
doc.add_paragraph('5. Priorisierte Maßnahmen-Roadmap', style='SectionHead')

doc.add_paragraph('Phase 1 – Sofort (1–2 Tage)', style='SubHead')
actions_p1 = [
    'SEC-001: Signup-Endpunkt absichern (Admin-Only oder deaktivieren)',
    'SEC-002: Rate-Limiting für Login/Signup implementieren',
    'SEC-003: Passwort-Komplexitätsregeln einführen',
]
for a in actions_p1:
    doc.add_paragraph(a, style='List Bullet')

doc.add_paragraph('Phase 2 – Kurzfristig (1 Woche)', style='SubHead')
actions_p2 = [
    'SEC-004: Zod-Validierungsschemas für alle Entitäten',
    'SEC-005: Security-HTTP-Header konfigurieren',
    'SEC-006: JWT-Ablaufzeit auf 8 Stunden reduzieren',
]
for a in actions_p2:
    doc.add_paragraph(a, style='List Bullet')

doc.add_paragraph('Phase 3 – Mittelfristig (2–4 Wochen)', style='SubHead')
actions_p3 = [
    'SEC-007: Umfassendes Audit-Logging implementieren',
    'SEC-008: Explizite CORS-Policy definieren',
    'SEC-009: Strukturiertes Logging mit Log-Leveln',
]
for a in actions_p3:
    doc.add_paragraph(a, style='List Bullet')

# ========== FOOTER ==========
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.style.font.size = Pt(8)
fp.style.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
run = fp.add_run('BANDspirit Sicherheitsaudit – VERTRAULICH – ')
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
# Page number
run2 = fp.add_run()
fldChar1 = run2._r.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'begin'})
run2._r.append(fldChar1)
run3 = fp.add_run()
instrText = run3._r.makeelement(qn('w:instrText'), {})
instrText.text = ' PAGE '
run3._r.append(instrText)
run4 = fp.add_run()
fldChar2 = run4._r.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'end'})
run4._r.append(fldChar2)

output_path = '/home/ubuntu/klientenmanagement/nextjs_space/public/BANDspirit_Sicherheitsaudit_v1.2.docx'
doc.save(output_path)
print(f'Report saved to {output_path}')
