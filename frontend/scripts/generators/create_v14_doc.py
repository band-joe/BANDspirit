from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
import datetime

doc = Document()

# Page setup
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

# Styles
style_title = doc.styles.add_style('DocTitle', WD_STYLE_TYPE.PARAGRAPH)
style_title.font.name = 'Calibri'
style_title.font.size = Pt(26)
style_title.font.bold = True
style_title.font.color.rgb = RGBColor(0x1a, 0x23, 0x7e)
style_title.paragraph_format.space_after = Pt(6)
style_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

style_subtitle = doc.styles.add_style('DocSubtitle', WD_STYLE_TYPE.PARAGRAPH)
style_subtitle.font.name = 'Calibri'
style_subtitle.font.size = Pt(14)
style_subtitle.font.color.rgb = RGBColor(0x0d, 0x47, 0xa1)
style_subtitle.paragraph_format.space_after = Pt(4)
style_subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

style_section = doc.styles.add_style('SectionHead', WD_STYLE_TYPE.PARAGRAPH)
style_section.font.name = 'Calibri'
style_section.font.size = Pt(16)
style_section.font.bold = True
style_section.font.color.rgb = RGBColor(0x1a, 0x23, 0x7e)
style_section.paragraph_format.space_before = Pt(18)
style_section.paragraph_format.space_after = Pt(8)

style_subsection = doc.styles.add_style('SubSectionHead', WD_STYLE_TYPE.PARAGRAPH)
style_subsection.font.name = 'Calibri'
style_subsection.font.size = Pt(13)
style_subsection.font.bold = True
style_subsection.font.color.rgb = RGBColor(0x0d, 0x47, 0xa1)
style_subsection.paragraph_format.space_before = Pt(12)
style_subsection.paragraph_format.space_after = Pt(6)

style_body = doc.styles['Normal']
style_body.font.name = 'Calibri'
style_body.font.size = Pt(11)
style_body.paragraph_format.space_after = Pt(6)
style_body.paragraph_format.line_spacing = Pt(16)

# Helper
def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(f' {text}')
    else:
        p.add_run(text)
    return p

def add_table_row(table, cells_data, bold=False, bg_color=None):
    row = table.add_row()
    for i, text in enumerate(cells_data):
        cell = row.cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        r = p.add_run(str(text))
        r.font.size = Pt(9.5)
        r.font.name = 'Calibri'
        if bold:
            r.bold = True
        if bg_color:
            from docx.oxml.ns import qn
            shading = cell._element.makeelement(qn('w:shd'), {
                qn('w:val'): 'clear',
                qn('w:color'): 'auto',
                qn('w:fill'): bg_color,
            })
            cell._element.get_or_add_tcPr = lambda: cell._element.find(qn('w:tcPr')) or cell._element.makeelement(qn('w:tcPr'), {})
    return row

def make_header_row(table, headers):
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.name = 'Calibri'

def add_data_row(table, data):
    row = table.add_row()
    for i, text in enumerate(data):
        cell = row.cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(str(text))
        r.font.size = Pt(9.5)
        r.font.name = 'Calibri'
    return row

# ============ COVER PAGE ============
p = doc.add_paragraph('BANDspirit', style='DocTitle')
p = doc.add_paragraph('Klientenmanagement-System', style='DocSubtitle')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('\n')
r = p.add_run('Content Engineering Dokument')
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1a, 0x23, 0x7e)
p.add_run('\n')
r2 = p.add_run('Version 1.4')
r2.font.size = Pt(14)
r2.font.color.rgb = RGBColor(0x0d, 0x47, 0xa1)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('\n\n')
r = p.add_run(f'Stand: {datetime.date.today().strftime("%d.%m.%Y")}')
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
p.add_run('\n')
r = p.add_run('bandkms.abacusai.app')
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0x0d, 0x47, 0xa1)

doc.add_page_break()

# ============ \u00c4NDERUNGSHISTORIE ============
doc.add_paragraph('\u00c4nderungshistorie', style='SectionHead')

table = doc.add_table(rows=1, cols=3, style='Table Grid')
table.alignment = WD_TABLE_ALIGNMENT.LEFT
make_header_row(table, ['Version', 'Datum', '\u00c4nderungen'])

changes = [
    ('1.4', '06.05.2026', 'Modul 4 Erweiterung: Einsatzplanung \u2014 Mutation und Kopie. Feldbearbeitung bestehender Einsatzpl\u00e4ne, Duplizierung als neue GEPLANT-Eins\u00e4tze.'),
    ('1.3', '05.05.2026', 'Modul 12: Meeting-Planer \u2014 Wochenkalender, ICS-Einladungen, Meeting-Er\u00f6ffnung'),
    ('1.2', '05.05.2026', 'Modul 11: Meetings, Transkription, KI-Protokoll; Security Hardening'),
    ('1.1', '05.05.2026', 'Modul 10: Benutzerverwaltung, User-Kontakt-Verkn\u00fcpfung'),
    ('1.0', '30.04.2026', 'Initiale Version, Module 1\u20139'),
]
for v, d, c in changes:
    add_data_row(table, [v, d, c])

# ============ MODULE OVERVIEW ============
doc.add_paragraph('Modul\u00fcbersicht', style='SectionHead')

table2 = doc.add_table(rows=1, cols=3, style='Table Grid')
make_header_row(table2, ['#', 'Modul', 'Beschreibung'])

modules = [
    ('1', 'Dashboard', 'KPI-Karten, Aktivit\u00e4ten, Workflows'),
    ('2', 'Klienten', 'Stammdaten CRUD, 360\u00b0-Ansicht'),
    ('3', 'Intake', 'Checklisten-Workflow'),
    ('4', 'Einsatzplanung', 'Listen-/Statusansicht, Mutation und Kopie (v1.4)'),
    ('5', 'Arbeitspl\u00e4tze', 'Kapazit\u00e4ts-Dashboard'),
    ('6', 'Berufsbilder', 'Katalog, Kompetenz-Tags'),
    ('7', 'Berichtswesen', 'Vorlagen, PDF-Export'),
    ('8', 'Abrechnungen', 'Freigabe-Workflow, DATEV-Export'),
    ('9', 'Kontakte', 'Kontaktpersonen, M:N-Zuordnung'),
    ('10', 'Benutzerverwaltung', 'CRUD, RBAC, Soft-Delete (v1.1)'),
    ('11', 'Meetings', 'Aufnahme, Transkription, KI-Protokoll (v1.2)'),
    ('12', 'Meeting-Planer', 'Wochenkalender, ICS-Einladungen, Meeting-Er\u00f6ffnung (v1.3)'),
]
for num, name, desc in modules:
    row = table2.add_row()
    for i, t in enumerate([num, name, desc]):
        cell = row.cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(t)
        r.font.size = Pt(9.5)
        if i == 1:
            r.bold = True

doc.add_page_break()

# ============ V1.4 CONTENT: EINSATZPLANUNG MUTATION & KOPIE ============
doc.add_paragraph('Modul 4 Erweiterung: Einsatzplanung \u2014 Mutation und Kopie (v1.4)', style='SectionHead')

doc.add_paragraph('\u00dcbersicht', style='SubSectionHead')
doc.add_paragraph(
    'Bestehende Einsatzpl\u00e4ne k\u00f6nnen nach der Erstellung mutiert (Feld\u00e4nderungen) und kopiert '
    '(Duplizierung als neuer Einsatzplan) werden. Die Erweiterung ber\u00fccksichtigt statusabh\u00e4ngige '
    'Einschr\u00e4nkungen, Gesch\u00e4ftsregel-Revalidierung und l\u00fcckenlose Audit-Protokollierung.'
)

doc.add_paragraph('Ist-Analyse', style='SubSectionHead')
doc.add_paragraph(
    'Der bestehende PUT-Endpunkt /api/einsatzplaene/[id] unterst\u00fctzt aktuell ausschlie\u00dflich '
    'Status\u00fcberg\u00e4nge (z.B. GEPLANT \u2192 AKTIV, AKTIV \u2192 BEENDET). Eine \u00c4nderung der '
    'Sachfelder (Arbeitsplatz, Stunden/Woche, Start-/Enddatum, Berufsbild) ist nach der Erstellung '
    'nicht m\u00f6glich. Ebenso fehlt eine Kopierfunktion zur Duplizierung bestehender Einsatzpl\u00e4ne.'
)

# --- FR-4.1 MUTATION ---
doc.add_paragraph('FR-4.1: Einsatzplan mutieren (Feldbearbeitung)', style='SubSectionHead')

doc.add_paragraph('Ein bestehender Einsatzplan kann in folgenden Feldern bearbeitet werden:')

tbl_fields = doc.add_table(rows=1, cols=3, style='Table Grid')
make_header_row(tbl_fields, ['Feld', 'Typ', 'Beschreibung'])
fields = [
    ('arbeitsplatzId', 'String', 'Zuordnung zu einem anderen Arbeitsplatz'),
    ('berufsbildId', 'String?', 'Zuordnung zu einem anderen/keinem Berufsbild'),
    ('stundenWoche', 'Decimal', 'W\u00f6chentliche Arbeitsstunden'),
    ('startDatum', 'DateTime', 'Beginn des Einsatzes'),
    ('endDatum', 'DateTime?', 'Ende des Einsatzes (optional)'),
]
for f, t, d in fields:
    add_data_row(tbl_fields, [f, t, d])

doc.add_paragraph('\nStatusabh\u00e4ngige Mutationsregeln:')

tbl_status = doc.add_table(rows=1, cols=3, style='Table Grid')
make_header_row(tbl_status, ['Status', 'Mutation erlaubt?', 'Einschr\u00e4nkungen'])
status_rules = [
    ('GEPLANT', '\u2713 Vollst\u00e4ndig', 'Alle Felder frei editierbar'),
    ('AKTIV', '\u2713 Eingeschr\u00e4nkt', 'Nur stundenWoche, endDatum, berufsbildId. arbeitsplatzId und startDatum gesperrt.'),
    ('PAUSIERT', '\u2713 Eingeschr\u00e4nkt', 'Wie AKTIV \u2014 nur stundenWoche, endDatum, berufsbildId \u00e4nderbar'),
    ('BEENDET', '\u2717 Gesperrt', 'Keine Mutation m\u00f6glich (abgeschlossener Einsatz)'),
]
for s, m, e in status_rules:
    add_data_row(tbl_status, [s, m, e])

doc.add_paragraph('\nGesch\u00e4ftsregeln bei Mutation:', style='SubSectionHead')

mutation_rules = [
    ('Kapazit\u00e4tspr\u00fcfung:', 'Bei \u00c4nderung von arbeitsplatzId wird die Kapazit\u00e4t des neuen Arbeitsplatzes gepr\u00fcft (aktive Eins\u00e4tze < maxKapazitaet). Die Kapazit\u00e4t des alten Arbeitsplatzes wird freigegeben.'),
    ('Datumsvalidierung:', 'endDatum muss nach startDatum liegen (falls gesetzt). startDatum darf bei GEPLANT-Status nicht in der Vergangenheit ge\u00e4ndert werden.'),
    ('Klient-Validierung:', 'klientId ist nicht mutierbar \u2014 ein Einsatz bleibt immer dem urspr\u00fcnglichen Klienten zugeordnet.'),
    ('Audit-Log:', 'Jede Feld\u00e4nderung wird im AuditLog mit Vor-/Nachwert protokolliert (Aktion: EINSATZ_MUTIERT).'),
]
for prefix, text in mutation_rules:
    add_bullet(text, prefix)

doc.add_paragraph('\nAPI-Erweiterung:')
doc.add_paragraph(
    'PUT /api/einsatzplaene/[id] \u2014 Erweitert um Feld-Mutation (zus\u00e4tzlich zum bestehenden '
    'Status-\u00dcbergang). Der Request-Body kann nun sowohl status als auch Feld\u00e4nderungen enthalten.'
)

# --- FR-4.2 KOPIE ---
doc.add_paragraph('FR-4.2: Einsatzplan kopieren (Duplizierung)', style='SubSectionHead')

doc.add_paragraph(
    'Ein bestehender Einsatzplan kann als Vorlage f\u00fcr einen neuen Einsatzplan dupliziert werden.'
)

doc.add_paragraph('Kopierverhalten:')
copy_items = [
    'Es wird ein neuer Einsatzplan erstellt mit Status GEPLANT',
    'Folgende Felder werden \u00fcbernommen: arbeitsplatzId, berufsbildId, stundenWoche',
    'Folgende Felder werden NICHT \u00fcbernommen: startDatum, endDatum (m\u00fcssen neu gesetzt werden), status (immer GEPLANT)',
    'klientId wird \u00fcbernommen \u2014 kann aber im Erstelldialog ge\u00e4ndert werden',
]
for item in copy_items:
    add_bullet(item)

doc.add_paragraph('\nGesch\u00e4ftsregeln bei Kopie:')
copy_rules = [
    ('Max. 1 aktiver Einsatz:', 'Die bestehende Regel wird beachtet. Da der kopierte Einsatz im Status GEPLANT startet, wird die Regel erst bei Aktivierung gepr\u00fcft.'),
    ('Kapazit\u00e4tspr\u00fcfung:', 'Wird beim Kopieren nicht gepr\u00fcft (erst bei Aktivierung relevant).'),
    ('Klient-Validierung:', 'Klient muss weiterhin AKTIV sein und einen abgeschlossenen Intake haben.'),
]
for prefix, text in copy_rules:
    add_bullet(text, prefix)

doc.add_paragraph('\nBenutzerinteraktion:')
ui_items = [
    'In der Einsatzplan-Detailansicht und Listenansicht wird ein \u201eKopieren\u201c-Button angezeigt',
    'Beim Klick \u00f6ffnet sich das Einsatzplan-Erstellformular, vorausgef\u00fcllt mit den Daten des Quell-Einsatzes',
    'Der Benutzer kann alle Felder vor dem Speichern anpassen',
    'Nach dem Speichern wird der neue Einsatzplan erstellt (der Quell-Einsatz bleibt unver\u00e4ndert)',
]
for item in ui_items:
    add_bullet(item)

doc.add_paragraph('\nAPI-Endpunkt (neu):')
doc.add_paragraph(
    'POST /api/einsatzplaene/[id]/kopie \u2014 Erstellt eine Kopie des Einsatzplans mit der angegebenen ID. '
    'Gibt die kopierbaren Felder zur\u00fcck, die im Erstelldialog vorausgef\u00fcllt werden.'
)

# --- FR-4.3 UI ---
doc.add_paragraph('FR-4.3: UI-Erweiterungen', style='SubSectionHead')
ui_reqs = [
    ('Einsatzplan-Detailseite:', 'Neue Buttons \u201eBearbeiten\u201c und \u201eKopieren\u201c (kontextsensitiv je nach Status und Berechtigung)'),
    ('Einsatzplan-Listenansicht:', 'Kontextmen\u00fc mit Optionen \u201eBearbeiten\u201c, \u201eKopieren\u201c'),
    ('Bearbeitungsmodus:', 'Inline-Bearbeitung oder Modal-Dialog mit den mutierbaren Feldern (statusabh\u00e4ngig)'),
    ('Validierungsfeedback:', 'Echtzeit-Validierung der Gesch\u00e4ftsregeln (Kapazit\u00e4t, Datumslogik) im Formular'),
    ('Audit-Anzeige:', 'Mutationshistorie in der Detailansicht sichtbar (wer hat wann was ge\u00e4ndert)'),
]
for prefix, text in ui_reqs:
    add_bullet(text, prefix)

# --- FR-4.4 BERECHTIGUNGEN ---
doc.add_paragraph('FR-4.4: Berechtigungen', style='SubSectionHead')
doc.add_paragraph(
    'Die Mutation und Kopie nutzen die bestehenden Berechtigungen einsatz:update (Mutation) '
    'und einsatz:create (Kopie). Keine neuen RBAC-Schl\u00fcssel erforderlich.'
)

tbl_perms = doc.add_table(rows=1, cols=3, style='Table Grid')
make_header_row(tbl_perms, ['Aktion', 'Berechtigung', 'Rollen'])
perm_rows = [
    ('Einsatz mutieren', 'einsatz:update', 'ADMIN, EINSATZPLANER, TEAMLEITUNG'),
    ('Einsatz kopieren', 'einsatz:create', 'ADMIN, EINSATZPLANER, TEAMLEITUNG'),
]
for a, b, r in perm_rows:
    add_data_row(tbl_perms, [a, b, r])

# --- API-ENDPUNKTE ---
doc.add_paragraph('API-Endpunkte (v1.4)', style='SubSectionHead')

tbl_api = doc.add_table(rows=1, cols=3, style='Table Grid')
make_header_row(tbl_api, ['Pfad', 'Methode', 'Beschreibung'])
api_rows = [
    ('/api/einsatzplaene/[id]', 'PUT', 'Erweitert \u2014 Feld-Mutation + Status-\u00dcbergang'),
    ('/api/einsatzplaene/[id]/kopie', 'POST', 'Neu \u2014 Einsatzplan kopieren (Vorlagedaten)'),
]
for p, m, d in api_rows:
    add_data_row(tbl_api, [p, m, d])

# --- GESCH\u00c4FTSREGELN ZUSAMMENFASSUNG ---
doc.add_paragraph('Gesch\u00e4ftsregeln (v1.4 \u2014 Zusammenfassung)', style='SubSectionHead')

rules_v14 = [
    'Einsatzplan-Mutation: Nur im Status GEPLANT vollst\u00e4ndig editierbar; AKTIV/PAUSIERT eingeschr\u00e4nkt; BEENDET gesperrt',
    'Einsatzplan-Mutation: Bei Arbeitsplatzwechsel Kapazit\u00e4ts-Revalidierung (alter Platz freigeben, neuer Platz pr\u00fcfen)',
    'Einsatzplan-Mutation: klientId ist unver\u00e4nderbar (Zuordnung zum Klienten permanent)',
    'Einsatzplan-Mutation: Jede \u00c4nderung wird im AuditLog protokolliert (Aktion: EINSATZ_MUTIERT)',
    'Einsatzplan-Kopie: Erstellt neuen Einsatz im Status GEPLANT',
    'Einsatzplan-Kopie: Start-/Enddatum m\u00fcssen neu vergeben werden',
    'Einsatzplan-Kopie: Max-1-aktiv-Regel wird erst bei Aktivierung gepr\u00fcft',
    'Einsatzplan-Kopie: Klient muss AKTIV sein mit abgeschlossenem Intake',
]
for rule in rules_v14:
    add_bullet(rule)

# ============ BISHERIGE MODULE (REFERENZ) ============
doc.add_page_break()
doc.add_paragraph('Bisherige Module (Referenz)', style='SectionHead')

doc.add_paragraph(
    'Die vollst\u00e4ndige Dokumentation der Module 1\u201312 ist in den vorherigen Versionen '
    '(v1.0\u2013v1.3) des Content Engineering Dokuments enthalten. Nachfolgend eine Kurz\u00fcbersicht:'
)

ref_items = [
    ('Module 1\u20139 (v1.0):', 'Dashboard, Klienten, Intake, Einsatzplanung, Arbeitspl\u00e4tze, Berufsbilder, Berichtswesen, Abrechnungen, Kontakte'),
    ('Modul 10 (v1.1):', 'Benutzerverwaltung mit RBAC, Soft-Delete, User-Kontakt-Verkn\u00fcpfung'),
    ('Modul 11 (v1.2):', 'Meetings mit konfigurierbaren Typen, Audioaufnahme, Echtzeit-Transkription, KI-Protokollerstellung'),
    ('Modul 12 (v1.3):', 'Meeting-Planer mit Wochenkalender, ICS-Einladungen per E-Mail, Meeting-Er\u00f6ffnung aus Termin'),
]
for prefix, text in ref_items:
    add_bullet(text, prefix)

# ============ TESTBENUTZER ============
doc.add_paragraph('Testbenutzer', style='SectionHead')
doc.add_paragraph('Alle Passw\u00f6rter: Test123!')

tbl_users = doc.add_table(rows=1, cols=3, style='Table Grid')
make_header_row(tbl_users, ['E-Mail', 'Rolle', 'Kontaktverkn\u00fcpfung'])
users = [
    ('admin@klientenmanagement.de', 'Admin', '\u2014'),
    ('fallmanager@klientenmanagement.de', 'Fallmanager', 'Thomas Berger (Betreuer)'),
    ('planer@klientenmanagement.de', 'Einsatzplaner', '\u2014'),
    ('team@klientenmanagement.de', 'Teamleitung', '\u2014'),
    ('buchhaltung@klientenmanagement.de', 'Buchhalter', '\u2014'),
]
for e, r, k in users:
    add_data_row(tbl_users, [e, r, k])

# ============ FOOTER ============
section = doc.sections[0]
footer = section.footer
footer_para = footer.paragraphs[0]
footer_para.text = 'BANDspirit \u2014 Content Engineering Dokument v1.4'
footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in footer_para.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

# Save
outpath = '/home/ubuntu/klientenmanagement/nextjs_space/public/BANDspirit_Content_Engineering_v1.4.docx'
doc.save(outpath)
print(f'Dokument gespeichert: {outpath}')
